from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from ai.parsing.models import ParsedInputItem


TEXT_COLUMNS = [
    "raw_text",
    "update",
    "field_update",
    "progress_update",
    "remarks",
    "description",
    "work_description",
    "activity",
]


def clean_value(
    value: Any,
) -> str:
    if value is None:
        return ""

    return str(value).strip()


def normalize_header(
    value: Any,
) -> str:
    return clean_value(
        value
    ).lower()


def find_value(
    row: dict,
    possible_columns: list[str],
):
    for column in possible_columns:
        value = row.get(column)

        if value is None:
            continue

        cleaned = clean_value(value)

        if cleaned:
            return cleaned

    return None


def build_row_text(
    row: dict,
) -> str:

    direct = find_value(
        row,
        TEXT_COLUMNS,
    )

    if direct:
        return direct

    parts = []

    for key, value in row.items():
        cleaned = clean_value(value)

        if cleaned:
            parts.append(
                f"{key}: {cleaned}"
            )

    return " | ".join(parts)


def parse_xlsx(
    file_bytes: bytes,
) -> list[ParsedInputItem]:

    workbook = load_workbook(
        BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    items: list[ParsedInputItem] = []

    for worksheet in workbook.worksheets:

        rows = worksheet.iter_rows(
            values_only=True
        )

        try:
            headers_raw = next(rows)
        except StopIteration:
            continue

        headers = [
            normalize_header(header)
            for header in headers_raw
        ]

        for row_number, values in enumerate(
            rows,
            start=2,
        ):
            row = {}

            for index, value in enumerate(
                values
            ):
                if index >= len(headers):
                    continue

                header = headers[index]

                if not header:
                    continue

                row[header] = value

            raw_text = build_row_text(
                row
            )

            if not raw_text:
                continue

            discipline = find_value(
                row,
                ["discipline"],
            )

            area = find_value(
                row,
                ["area", "location"],
            )

            unit = find_value(
                row,
                ["unit", "uom"],
            )

            quantity_raw = find_value(
                row,
                [
                    "quantity",
                    "qty",
                    "actual_quantity",
                ],
            )

            quantity = None

            if quantity_raw is not None:
                try:
                    quantity = float(
                        quantity_raw
                    )
                except ValueError:
                    quantity = None

            items.append(
                ParsedInputItem(
                    raw_text=raw_text,
                    discipline=discipline,
                    area=area,
                    quantity=quantity,
                    unit=unit,
                    source_row=row_number,
                )
            )

    workbook.close()

    return items