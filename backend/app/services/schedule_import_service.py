from __future__ import annotations

import csv
import io
import re

from datetime import (
    date,
    datetime,
    timezone,
)

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.database.supabase import supabase

from app.services.xer_schedule_parser import (
    parse_xer_schedule,
)


SUPPORTED_EXTENSIONS = {
    ".csv",
    ".xlsx",
    ".xer",
}

INSERT_BATCH_SIZE = 200


COLUMN_ALIASES = {
    "activity_code": [
        "activity_code",
        "activity code",
        "activity id",
        "activity_id",
        "task id",
        "task_id",
        "id",
    ],
    "activity_name": [
        "activity_name",
        "activity name",
        "task name",
        "task_name",
        "name",
        "description",
    ],
    "wbs_level": [
        "wbs_level",
        "wbs level",
        "wbs",
        "wbs code",
        "wbs_code",
    ],
    "discipline": [
        "discipline",
        "trade",
    ],
    "area": [
        "area",
        "location",
        "zone",
    ],
    "planned_start": [
        "planned_start",
        "planned start",
        "start",
        "start date",
        "start_date",
        "baseline start",
        "baseline_start",
    ],
    "planned_finish": [
        "planned_finish",
        "planned finish",
        "finish",
        "finish date",
        "finish_date",
        "end",
        "end date",
        "end_date",
        "baseline finish",
        "baseline_finish",
    ],
    "planned_progress": [
        "planned_progress",
        "planned progress",
        "planned %",
        "planned percent",
        "planned percentage",
    ],
    "planned_quantity": [
        "planned_quantity",
        "planned quantity",
        "budgeted quantity",
        "budget quantity",
        "quantity",
        "qty",
    ],
    "quantity_unit": [
        "quantity_unit",
        "quantity unit",
        "unit",
        "uom",
    ],
    "predecessors": [
        "predecessors",
        "predecessor",
        "predecessor activities",
        "predecessor activity",
        "predecessor ids",
        "predecessor id",
        "pred",
    ],
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""

    return " ".join(
        str(value)
        .strip()
        .split()
    )


def normalize_header(value: Any) -> str:
    return (
        clean_text(value)
        .lower()
        .replace("-", " ")
        .replace("_", " ")
    )


def get_extension(filename: str) -> str:
    lower = filename.lower()

    for extension in (
        ".xlsx",
        ".csv",
        ".xer",
    ):
        if lower.endswith(extension):
            return extension

    return ""


def chunk_list(
    values: list,
    size: int,
):
    for start in range(
        0,
        len(values),
        size,
    ):
        yield values[
            start:
            start + size
        ]


def find_column_value(
    row: dict[str, Any],
    logical_name: str,
):
    aliases = COLUMN_ALIASES.get(
        logical_name,
        [],
    )

    normalized_row = {
        normalize_header(key): value
        for key, value
        in row.items()
    }

    for alias in aliases:
        normalized_alias = normalize_header(
            alias
        )

        if normalized_alias not in normalized_row:
            continue

        value = normalized_row[
            normalized_alias
        ]

        if value is None:
            continue

        if clean_text(value):
            return value

    return None


def parse_float(
    value: Any,
) -> float | None:
    if value is None:
        return None

    if isinstance(
        value,
        (
            int,
            float,
        ),
    ):
        return float(value)

    text = (
        clean_text(value)
        .replace(",", "")
        .replace("%", "")
    )

    if not text:
        return None

    try:
        return float(text)
    except ValueError:
        return None


DATE_FORMATS = [
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%d-%m-%Y",
    "%d/%m/%Y",
    "%m/%d/%Y",
    "%d-%b-%Y",
    "%d %b %Y",
    "%d-%B-%Y",
    "%d %B %Y",
]


def parse_datetime_value(
    value: Any,
) -> str | None:
    if value is None:
        return None

    if isinstance(
        value,
        datetime,
    ):
        parsed = value

        if parsed.tzinfo is None:
            parsed = parsed.replace(
                tzinfo=timezone.utc
            )

        return parsed.isoformat()

    if isinstance(
        value,
        date,
    ):
        parsed = datetime(
            value.year,
            value.month,
            value.day,
            tzinfo=timezone.utc,
        )

        return parsed.isoformat()

    text = clean_text(value)

    if not text:
        return None

    try:
        parsed = datetime.fromisoformat(
            text.replace(
                "Z",
                "+00:00",
            )
        )

        if parsed.tzinfo is None:
            parsed = parsed.replace(
                tzinfo=timezone.utc
            )

        return parsed.isoformat()

    except ValueError:
        pass

    for date_format in DATE_FORMATS:
        try:
            parsed = (
                datetime.strptime(
                    text,
                    date_format,
                )
                .replace(
                    tzinfo=timezone.utc
                )
            )

            return parsed.isoformat()

        except ValueError:
            continue

    return None


def parse_csv_rows(
    file_bytes: bytes,
) -> list[dict[str, Any]]:
    text = file_bytes.decode(
        "utf-8-sig",
        errors="replace",
    )

    reader = csv.DictReader(
        io.StringIO(text)
    )

    rows = []

    for row_number, row in enumerate(
        reader,
        start=2,
    ):
        normalized = dict(row)
        normalized[
            "__source_row"
        ] = row_number
        rows.append(normalized)

    return rows


def parse_xlsx_rows(
    file_bytes: bytes,
) -> list[dict[str, Any]]:
    workbook = load_workbook(
        BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    rows = []

    for worksheet in workbook.worksheets:
        worksheet_rows = worksheet.iter_rows(
            values_only=True
        )

        try:
            headers = next(
                worksheet_rows
            )
        except StopIteration:
            continue

        headers = [
            clean_text(header)
            for header in headers
        ]

        for row_number, values in enumerate(
            worksheet_rows,
            start=2,
        ):
            row = {
                "__source_row":
                    row_number,
                "__sheet":
                    worksheet.title,
            }

            has_data = False

            for index, value in enumerate(values):
                if index >= len(headers):
                    continue

                header = headers[index]

                if not header:
                    continue

                row[header] = value

                if value not in (
                    None,
                    "",
                ):
                    has_data = True

            if has_data:
                rows.append(row)

    workbook.close()
    return rows


def parse_raw_rows(
    filename: str,
    file_bytes: bytes,
) -> list[dict[str, Any]]:
    extension = get_extension(
        filename
    )

    if extension not in (
        ".csv",
        ".xlsx",
    ):
        raise ValueError(
            "Structured row parsing supports CSV and XLSX."
        )

    if extension == ".csv":
        return parse_csv_rows(
            file_bytes
        )

    return parse_xlsx_rows(
        file_bytes
    )


def parse_predecessor_token(
    token: str,
) -> dict | None:
    token = clean_text(token)

    if not token:
        return None

    token = token.replace(
        " ",
        "",
    )

    pattern = re.compile(
        r"^(?P<code>.+?)"
        r"(?P<type>FS|SS|FF|SF)?"
        r"(?P<lag>[+-]\d+(?:\.\d+)?)?"
        r"(?:d|day|days)?$",
        re.IGNORECASE,
    )

    match = pattern.match(token)

    if not match:
        return {
            "predecessor_code":
                token,
            "dependency_type":
                "FS",
            "lag_days":
                0,
        }

    code = (
        match.group("code")
        or ""
    ).strip()

    dependency_type = (
        match.group("type")
        or "FS"
    ).upper()

    lag_raw = (
        match.group("lag")
        or "0"
    )

    try:
        lag_days = int(
            round(
                float(lag_raw)
            )
        )
    except ValueError:
        lag_days = 0

    return {
        "predecessor_code":
            code,
        "dependency_type":
            dependency_type,
        "lag_days":
            lag_days,
    }


def parse_predecessors(
    value: Any,
) -> list[dict]:
    text = clean_text(value)

    if not text:
        return []

    tokens = re.split(
        r"[,;\n]+",
        text,
    )

    dependencies = []

    for token in tokens:
        dependency = parse_predecessor_token(
            token
        )

        if dependency:
            dependencies.append(
                dependency
            )

    return dependencies


def normalize_schedule_rows(
    rows: list[dict[str, Any]],
) -> dict:
    activities = []
    dependencies = []

    errors = []
    warnings = []

    seen_codes = set()

    for row in rows:
        source_row = row.get(
            "__source_row"
        )

        activity_code = clean_text(
            find_column_value(
                row,
                "activity_code",
            )
        )

        activity_name = clean_text(
            find_column_value(
                row,
                "activity_name",
            )
        )

        if (
            not activity_code
            and not activity_name
        ):
            continue

        if not activity_code:
            errors.append({
                "row":
                    source_row,
                "field":
                    "activity_code",
                "message":
                    "Activity code is required.",
            })
            continue

        if not activity_name:
            errors.append({
                "row":
                    source_row,
                "field":
                    "activity_name",
                "message":
                    f"Activity name is required for {activity_code}.",
            })
            continue

        normalized_code = activity_code.upper()

        if normalized_code in seen_codes:
            errors.append({
                "row":
                    source_row,
                "field":
                    "activity_code",
                "message":
                    f"Duplicate activity code {activity_code}.",
            })
            continue

        seen_codes.add(
            normalized_code
        )

        planned_start_raw = find_column_value(
            row,
            "planned_start",
        )

        planned_finish_raw = find_column_value(
            row,
            "planned_finish",
        )

        planned_start = parse_datetime_value(
            planned_start_raw
        )

        planned_finish = parse_datetime_value(
            planned_finish_raw
        )

        if (
            planned_start_raw
            and not planned_start
        ):
            errors.append({
                "row":
                    source_row,
                "field":
                    "planned_start",
                "message":
                    f"Invalid planned start date for {activity_code}.",
            })

        if (
            planned_finish_raw
            and not planned_finish
        ):
            errors.append({
                "row":
                    source_row,
                "field":
                    "planned_finish",
                "message":
                    f"Invalid planned finish date for {activity_code}.",
            })

        planned_progress = parse_float(
            find_column_value(
                row,
                "planned_progress",
            )
        )

        if planned_progress is None:
            planned_progress = 0.0

        planned_progress = max(
            0,
            min(
                100,
                planned_progress,
            ),
        )

        activity = {
            "activity_code":
                activity_code,
            "activity_name":
                activity_name,
            "wbs_level":
                clean_text(
                    find_column_value(
                        row,
                        "wbs_level",
                    )
                )
                or None,
            "discipline":
                clean_text(
                    find_column_value(
                        row,
                        "discipline",
                    )
                )
                or None,
            "area":
                clean_text(
                    find_column_value(
                        row,
                        "area",
                    )
                )
                or None,
            "planned_start":
                planned_start,
            "planned_finish":
                planned_finish,
            "planned_progress":
                planned_progress,

            # CSV/XLSX schedule imports do not currently claim
            # actual progress unless a dedicated actual-progress
            # column is added later.
            "actual_progress":
                0.0,

            "planned_quantity":
                parse_float(
                    find_column_value(
                        row,
                        "planned_quantity",
                    )
                ),
            "quantity_unit":
                clean_text(
                    find_column_value(
                        row,
                        "quantity_unit",
                    )
                )
                or None,
            "status":
                "not_started",
            "source_row":
                source_row,
        }

        activities.append(activity)

        predecessor_value = find_column_value(
            row,
            "predecessors",
        )

        for predecessor in parse_predecessors(
            predecessor_value
        ):
            dependencies.append({
                "successor_code":
                    activity_code,
                "predecessor_code":
                    predecessor[
                        "predecessor_code"
                    ],
                "dependency_type":
                    predecessor[
                        "dependency_type"
                    ],
                "lag_days":
                    predecessor[
                        "lag_days"
                    ],
                "source_row":
                    source_row,
            })

    available_codes = {
        activity[
            "activity_code"
        ].upper()
        for activity in activities
    }

    for dependency in dependencies:
        predecessor_code = dependency[
            "predecessor_code"
        ].upper()

        successor_code = dependency[
            "successor_code"
        ].upper()

        if predecessor_code not in available_codes:
            warnings.append({
                "row":
                    dependency[
                        "source_row"
                    ],
                "field":
                    "predecessors",
                "message":
                    (
                        "Predecessor "
                        f"{dependency['predecessor_code']} "
                        "is not present in the uploaded file."
                    ),
            })

        if predecessor_code == successor_code:
            errors.append({
                "row":
                    dependency[
                        "source_row"
                    ],
                "field":
                    "predecessors",
                "message":
                    "An activity cannot depend on itself.",
            })

    return {
        "activities":
            activities,
        "dependencies":
            dependencies,
        "errors":
            errors,
        "warnings":
            warnings,
        "source_rows":
            len(rows),
    }


def parse_schedule(
    filename: str,
    file_bytes: bytes,
) -> dict:
    extension = get_extension(
        filename
    )

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported schedule file. Supported formats: CSV, XLSX and XER."
        )

    if extension == ".xer":
        return parse_xer_schedule(
            file_bytes
        )

    rows = parse_raw_rows(
        filename,
        file_bytes,
    )

    return normalize_schedule_rows(
        rows
    )


def get_project(
    project_id: str,
) -> dict | None:
    response = (
        supabase
        .table("projects")
        .select(
            "id,name,start_date,end_date,status"
        )
        .eq(
            "id",
            project_id,
        )
        .limit(1)
        .execute()
    )

    if not response.data:
        return None

    return response.data[0]


def get_all_project_activities(
    project_id: str,
) -> list[dict]:
    all_rows = []

    page_size = 1000
    start = 0

    while True:
        response = (
            supabase
            .table("activities")
            .select(
                "id,activity_code"
            )
            .eq(
                "project_id",
                project_id,
            )
            .range(
                start,
                start + page_size - 1,
            )
            .execute()
        )

        rows = response.data or []
        all_rows.extend(rows)

        if len(rows) < page_size:
            break

        start += page_size

    return all_rows


def get_all_dependencies(
    project_id: str,
) -> list[dict]:
    all_rows = []

    page_size = 1000
    start = 0

    while True:
        response = (
            supabase
            .table(
                "activity_dependencies"
            )
            .select(
                "predecessor_activity_id,"
                "successor_activity_id,"
                "dependency_type,"
                "lag_days"
            )
            .eq(
                "project_id",
                project_id,
            )
            .range(
                start,
                start + page_size - 1,
            )
            .execute()
        )

        rows = response.data or []
        all_rows.extend(rows)

        if len(rows) < page_size:
            break

        start += page_size

    return all_rows


def preview_schedule_import(
    project_id: str,
    filename: str,
    file_bytes: bytes,
) -> dict:
    project = get_project(
        project_id
    )

    if not project:
        raise ValueError(
            "Project not found."
        )

    normalized = parse_schedule(
        filename,
        file_bytes,
    )

    existing_activities = get_all_project_activities(
        project_id
    )

    existing_codes = {
        row[
            "activity_code"
        ].upper()
        for row in existing_activities
        if row.get(
            "activity_code"
        )
    }

    new_count = 0
    existing_count = 0

    for activity in normalized[
        "activities"
    ]:
        if (
            activity[
                "activity_code"
            ].upper()
            in existing_codes
        ):
            existing_count += 1
        else:
            new_count += 1

    return {
        "project":
            project,
        "filename":
            filename,
        "valid":
            len(
                normalized[
                    "errors"
                ]
            ) == 0,
        "summary": {
            "source_rows":
                normalized.get(
                    "source_rows",
                    len(
                        normalized[
                            "activities"
                        ]
                    ),
                ),
            "activities":
                len(
                    normalized[
                        "activities"
                    ]
                ),
            "new_activities":
                new_count,
            "existing_activities":
                existing_count,
            "dependencies":
                len(
                    normalized[
                        "dependencies"
                    ]
                ),
            "errors":
                len(
                    normalized[
                        "errors"
                    ]
                ),
            "warnings":
                len(
                    normalized[
                        "warnings"
                    ]
                ),
        },
        "activities_preview":
            normalized[
                "activities"
            ][:50],
        "dependencies_preview":
            normalized[
                "dependencies"
            ][:50],
        "errors":
            normalized[
                "errors"
            ],
        "warnings":
            normalized[
                "warnings"
            ],
    }


def import_schedule(
    project_id: str,
    filename: str,
    file_bytes: bytes,
) -> dict:
    project = get_project(
        project_id
    )

    if not project:
        raise ValueError(
            "Project not found."
        )

    normalized = parse_schedule(
        filename,
        file_bytes,
    )

    if normalized[
        "errors"
    ]:
        raise ValueError(
            "Schedule contains validation errors. Preview and fix the file before importing."
        )

    existing_rows = get_all_project_activities(
        project_id
    )

    existing_by_code = {
        row[
            "activity_code"
        ].upper():
            row
        for row in existing_rows
        if row.get(
            "activity_code"
        )
    }

    new_activities = []
    skipped_existing = []

    for activity in normalized[
        "activities"
    ]:
        code_upper = activity[
            "activity_code"
        ].upper()

        if code_upper in existing_by_code:
            skipped_existing.append(
                activity[
                    "activity_code"
                ]
            )
            continue

        actual_progress = max(
            0.0,
            min(
                100.0,
                float(
                    activity.get(
                        "actual_progress",
                        0,
                    )
                    or 0
                ),
            ),
        )

        status = (
            activity.get(
                "status"
            )
            or "not_started"
        )

        if status == "completed":
            actual_progress = 100.0

        if (
            status == "not_started"
            and actual_progress > 0
        ):
            status = "in_progress"

        new_activities.append({
            "project_id":
                project_id,
            "activity_code":
                activity[
                    "activity_code"
                ],
            "activity_name":
                activity[
                    "activity_name"
                ],
            "wbs_level":
                activity.get(
                    "wbs_level"
                ),
            "discipline":
                activity.get(
                    "discipline"
                ),
            "area":
                activity.get(
                    "area"
                ),
            "planned_start":
                activity.get(
                    "planned_start"
                ),
            "planned_finish":
                activity.get(
                    "planned_finish"
                ),
            "planned_progress":
                activity.get(
                    "planned_progress",
                    0,
                ),
            "actual_progress":
                actual_progress,
            "status":
                status,
            "planned_quantity":
                activity.get(
                    "planned_quantity"
                ),
            "quantity_unit":
                activity.get(
                    "quantity_unit"
                ),
            "actual_quantity":
                0,
        })

    inserted_activity_rows = []

    for batch in chunk_list(
        new_activities,
        INSERT_BATCH_SIZE,
    ):
        response = (
            supabase
            .table(
                "activities"
            )
            .insert(
                batch
            )
            .execute()
        )

        inserted_activity_rows.extend(
            response.data
            or []
        )

    all_project_activities = get_all_project_activities(
        project_id
    )

    activity_id_by_code = {
        row[
            "activity_code"
        ].upper():
            row[
                "id"
            ]
        for row in all_project_activities
        if (
            row.get(
                "activity_code"
            )
            and row.get(
                "id"
            )
        )
    }

    existing_dependencies = get_all_dependencies(
        project_id
    )

    existing_dependency_keys = {
        (
            row.get(
                "predecessor_activity_id"
            ),
            row.get(
                "successor_activity_id"
            ),
            (
                row.get(
                    "dependency_type"
                )
                or "FS"
            ).upper(),
            int(
                row.get(
                    "lag_days"
                )
                or 0
            ),
        )
        for row in existing_dependencies
    }

    dependency_payloads = []
    skipped_dependencies = 0
    unresolved_dependencies = []

    for dependency in normalized[
        "dependencies"
    ]:
        predecessor_id = (
            activity_id_by_code.get(
                dependency[
                    "predecessor_code"
                ].upper()
            )
        )

        successor_id = (
            activity_id_by_code.get(
                dependency[
                    "successor_code"
                ].upper()
            )
        )

        if (
            not predecessor_id
            or not successor_id
        ):
            unresolved_dependencies.append(
                dependency
            )
            continue

        dependency_type = (
            dependency.get(
                "dependency_type"
            )
            or "FS"
        ).upper()

        lag_days = int(
            dependency.get(
                "lag_days"
            )
            or 0
        )

        dependency_key = (
            predecessor_id,
            successor_id,
            dependency_type,
            lag_days,
        )

        if dependency_key in existing_dependency_keys:
            skipped_dependencies += 1
            continue

        existing_dependency_keys.add(
            dependency_key
        )

        dependency_payloads.append({
            "project_id":
                project_id,
            "predecessor_activity_id":
                predecessor_id,
            "successor_activity_id":
                successor_id,
            "dependency_type":
                dependency_type,
            "lag_days":
                lag_days,
        })

    inserted_dependency_rows = []

    for batch in chunk_list(
        dependency_payloads,
        INSERT_BATCH_SIZE,
    ):
        response = (
            supabase
            .table(
                "activity_dependencies"
            )
            .insert(
                batch
            )
            .execute()
        )

        inserted_dependency_rows.extend(
            response.data
            or []
        )

    return {
        "success":
            True,
        "project":
            project,
        "filename":
            filename,
        "summary": {
            "activities_in_file":
                len(
                    normalized[
                        "activities"
                    ]
                ),
            "activities_inserted":
                len(
                    inserted_activity_rows
                ),
            "activities_skipped_existing":
                len(
                    skipped_existing
                ),
            "dependencies_in_file":
                len(
                    normalized[
                        "dependencies"
                    ]
                ),
            "dependencies_inserted":
                len(
                    inserted_dependency_rows
                ),
            "dependencies_skipped_existing":
                skipped_dependencies,
            "dependencies_unresolved":
                len(
                    unresolved_dependencies
                ),
        },
        "warnings":
            normalized[
                "warnings"
            ],
        "unresolved_dependencies":
            unresolved_dependencies[
                :50
            ],
    }

def clear_project_schedule(
    project_id: str,
) -> dict:
    """
    Clear a project's schedule only when it is safe to do so.

    We intentionally block deletion when field events or exception
    records already exist for the project, because those records may
    represent execution/history linked to schedule activities.
    """
    project = get_project(
        project_id
    )

    if not project:
        raise ValueError(
            "Project not found."
        )

    activities = get_all_project_activities(
        project_id
    )

    dependencies = get_all_dependencies(
        project_id
    )

    if not activities:
        return {
            "success": True,
            "project": project,
            "summary": {
                "activities_deleted": 0,
                "dependencies_deleted": 0,
            },
        }

    field_event_check = (
        supabase
        .table("field_events")
        .select("id")
        .eq(
            "project_id",
            project_id,
        )
        .limit(1)
        .execute()
    )

    if field_event_check.data:
        raise ValueError(
            "This schedule cannot be cleared because the project already "
            "contains field execution history. Create a fresh test project "
            "or remove the dependent test data first."
        )

    exception_check = (
        supabase
        .table("exceptions")
        .select("id")
        .eq(
            "project_id",
            project_id,
        )
        .limit(1)
        .execute()
    )

    if exception_check.data:
        raise ValueError(
            "This schedule cannot be cleared because project exceptions "
            "already reference the schedule. Clear the dependent test "
            "records first or use a fresh test project."
        )

    try:
        (
            supabase
            .table("activity_dependencies")
            .delete()
            .eq(
                "project_id",
                project_id,
            )
            .execute()
        )

        (
            supabase
            .table("activities")
            .delete()
            .eq(
                "project_id",
                project_id,
            )
            .execute()
        )

    except Exception as exc:
        raise ValueError(
            "The schedule could not be cleared because other project data "
            f"still references these activities: {str(exc)}"
        ) from exc

    return {
        "success": True,
        "project": project,
        "summary": {
            "activities_deleted":
                len(activities),
            "dependencies_deleted":
                len(dependencies),
        },
    }
